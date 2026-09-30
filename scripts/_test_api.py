import os
import json, urllib.request, sys
API = os.environ.get("WE_ENGINE","http://127.0.0.1:8020")
def post(path, payload):
    req = urllib.request.Request(API + path,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=60) as r:
        return json.loads(r.read().decode("utf-8"))

# 1. 建项目
r = post("/api/pert/project", {"name": "PERT端到端测试", "estimator": "李四"})
print("项目:", r)
pid = r["project_id"]

# 2. 增明细
ids=[]
d=[("需求分析","","","需求梳理",5,8,12),
   ("开发","票据系统","扣款","批量扣款提交",6,6,7),
   ("开发","票据系统","止付","止付结清",8,10,12),
   ("实施","","","部署调测",30,30,35)]
for ph,s,mod,desc,o,m,p in d:
    rr=post("/api/pert/item",{"project_id":pid,"phase":ph,"system":s,"module":mod,"desc":desc,
             "optimistic":o,"most_likely":m,"pessimistic":p})
    ids.append(rr["id"])
print("明细ids:", ids)

# 3. 列表+汇总
lst=post("/api/pert/list",{"project_id":pid})
print("汇总:", lst["summary"])

# 4. 生成
g=post("/api/pert/generate",{"project_id":pid,"est_date":"2026-09-08"})
print("生成:", g["ok"], g["sid"])

# 5. 验证文件
import openpyxl, urllib.request
dl = urllib.request.urlopen(API + "/api/download?s=" + g["sid"])
open(r"d:\code\otherProjects\27_Evaluation\data\generated\\" + g["sid"], "wb").write(dl.read())
wb = openpyxl.load_workbook(r"d:\code\otherProjects\27_Evaluation\data\generated\\" + g["sid"], data_only=False)
ws = wb[wb.sheetnames[0]]
print("D3(项目名):", ws["D3"].value, "| H2:", ws["H2"].value, "| D2:", ws["D2"].value)
print("开发R8 D/E/F/G:", ws.cell(8,4).value, ws.cell(8,5).value, ws.cell(8,6).value, ws.cell(8,7).value)
print("需求R6 E/F/G:", ws.cell(6,5).value, ws.cell(6,6).value, ws.cell(6,7).value)
print("实施R32 E/F/G:", ws.cell(32,5).value, ws.cell(32,6).value, ws.cell(32,7).value)
print("H8公式:", str(ws.cell(8,8).value)[:60])