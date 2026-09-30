# -*- coding: utf-8 -*-
"""定位 verify 差异：逐行对比 引擎US 与 母版M列(缓存值)。"""
import openpyxl
import engine
TP = r"C:\Users\hspcadmin\Desktop\工作量评估\副本造价评估-天津银行_票据新增五篇大文章标识功能项目（软件开发功能点评估）_KX_20251027_V1.0.xlsx"
wv = openpyxl.load_workbook(TP, data_only=True)["规模估算"]  # 缓存值
ws = openpyxl.load_workbook(TP, data_only=False)["规模估算"]  # 公式/原始
diff = 0
for r in range(6, 54):
    typ = ws.cell(r, 9).value
    desc = ws.cell(r, 7).value
    if not typ or desc is None:
        continue
    reuse = (ws.cell(r, 11).value or "").strip()
    modify = (ws.cell(r, 12).value or "").strip()
    m_val = wv.cell(r, 13).value
    try:
        eus = engine.us(str(typ).strip(), reuse, modify)
    except Exception as ex:
        print(f"r{r} EXC {ex} typ={typ!r} reuse={reuse!r} modify={modify!r}")
        continue
    flag = "  <<< DIFF" if abs((m_val or 0) - eus) > 0.005 else ""
    if flag:
        diff += 1
    print(f"r{r} {str(typ).strip():4} {reuse:2} {modify:2} 母版M={m_val} 引擎={eus}{flag} | {(desc or '')[:22]}")
print("差异行数:", diff)