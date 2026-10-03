# -*- coding: utf-8 -*-
"""PERT 三点估算：SQLite 数据层 + 计算引擎 + 阜新 V2.0 模板生成。
所有自包含，不依赖 server.py；frozen(exe) 时一并收集。
"""
import json, os, shutil, sqlite3
from pathlib import Path

# 阜新 V2.0 母版路径：环境 WE_TPL_PERT > 桌面默认
TPL_PERT = os.environ.get("WE_TPL_PERT") or \
    r"C:\Users\hspcadmin\Desktop\工作量评估\阜新银行-票据系统优化-工作量评估V2.0.xlsx"

# 六阶段固定顺序（对齐模板 R6/R7/R8/R32/R33/R34）
PHASES = ["需求分析", "系统设计", "开发", "实施", "公共及远程", "现场维护"]
PHASE_ROW = {"需求分析": 6, "系统设计": 7, "开发": 8, "实施": 32, "公共及远程": 33, "现场维护": 34}
DEV_START, DEV_END = 8, 30          # 开发明细数据区行号
SHEET = "项目工作量估算表"


class PertDB:
    """围绕单库 sqlite3 的轻封装，避免重复 connect。"""

    def __init__(self, db_path):
        self.db = db_path
        self._init()

    def _connect(self):
        # T01481（P1-1）：busy timeout + WAL，与服务端并发写一致
        con = sqlite3.connect(self.db, timeout=15)
        con.execute("PRAGMA journal_mode=WAL")
        con.execute("PRAGMA synchronous=NORMAL")
        return con

    def _init(self):
        con = self._connect()
        con.executescript("""
        CREATE TABLE IF NOT EXISTS projects(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          name TEXT NOT NULL,
          estimator TEXT DEFAULT '',
          created_at TEXT DEFAULT (datetime('now','localtime'))
        );
        CREATE TABLE IF NOT EXISTS pert_items(
          id INTEGER PRIMARY KEY AUTOINCREMENT,
          project_id INTEGER NOT NULL,
          phase TEXT NOT NULL DEFAULT '开发',
          system TEXT DEFAULT '', module TEXT DEFAULT '', desc TEXT DEFAULT '',
          optimistic REAL, most_likely REAL, pessimistic REAL
        );
        CREATE TABLE IF NOT EXISTS estimate_configs(
          project_id INTEGER PRIMARY KEY,
          cfg TEXT DEFAULT '{}'
        );
        """)
        con.commit()
        con.close()


# ---------------- 计算 ----------------
def exp(o, m, p):
    """PERT 预期 = (O + 4M + P) / 6。"""
    return (o + 4 * m + p) / 6


def sigma(o, p):
    """标准差 = (P - O) / 6。"""
    return (p - o) / 6


DEFAULT_CFG = {
    "days_per_man_month": 20.83,   # 每人月工天标准（名称管理器宏变量）
    "hours_per_day": 8,
    "adjust": {"客户协作": 0.06, "质量要求": 0.06, "范围或需求变更": 0.05, "进度变更": 0.06},
}


def load_cfg(con, pid):
    cfg = dict(DEFAULT_CFG)
    row = con.execute("SELECT cfg FROM estimate_configs WHERE project_id=?", (pid,)).fetchone()
    if row and row[0]:
        try:
            cfg.update(json.loads(row[0]))
        except Exception:
            pass
    return cfg


def save_cfg(con, pid, cfg):
    con.execute("INSERT OR REPLACE INTO estimate_configs(project_id,cfg) VALUES(?,?)",
                (pid, json.dumps(cfg, ensure_ascii=False)))
    con.commit()


def _phase_days(rows):
    """按阶段聚合预期人天与条目数，返回 (by_phase, total_days)。

    T01501（S3776）：原循环里同时做 None→float 兜底、期望值累加、字典初始化与
    嵌套字典自增，认知复杂度超阈值。拆出后 summarize 只剩「聚合 → 按阶段展开 → 换算人月」
    三步线性流程，聚合口径（浮点兜底为 0、total 累加顺序）完全不变。
    """
    by = {}
    total = 0.0
    for ph, o, m, p in rows:
        e = exp(float(o or 0), float(m or 0), float(p or 0))
        total += e
        slot = by.setdefault(ph, {"days": 0.0, "count": 0})
        slot["days"] += e
        slot["count"] += 1
    return by, total


def summarize(con, pid):
    """按阶段汇总预期人天与占比，叠加调节系数得人月。"""
    rows = con.execute(
        "SELECT phase,optimistic,most_likely,pessimistic FROM pert_items "
        "WHERE project_id=? ORDER BY id", (pid,)).fetchall()
    by, total = _phase_days(rows)
    phases = []
    for ph in PHASES:
        b = by.get(ph, {"days": 0.0, "count": 0})
        phases.append({"phase": ph, "days": round(b["days"], 2), "count": b["count"],
                       "pct": round(b["days"] / total * 100, 2) if total else 0})
    cfg = load_cfg(con, pid)
    dpm = float(cfg.get("days_per_man_month", 20.83))
    ideal_man = total / dpm
    adj_sum = sum(float(v) for v in cfg.get("adjust", {}).values())
    return {
        "total_days": round(total, 2),
        "ideal_person_month": round(ideal_man, 2),
        "adjust_sum": round(adj_sum, 4),
        "actual_person_month": round(ideal_man * (1 + adj_sum), 2),
        "phases": phases, "cfg": cfg,
    }


# ---------------- 阜新 V2.0 生成 ----------------
def generate(project_name, rows, out, est_date=None, estimator="", tpl=None):
    """套模板生成成品。rows: [{phase,system,module,desc,optimistic,most_likely,pessimistic}]
    开发明细写 R8 起(E/F/G)，其余阶段写阶段行(E/F/G)；H/I 列保留模板公式自动计算。
    """
    import openpyxl
    path = tpl or TPL_PERT
    if not Path(path).exists():
        raise FileNotFoundError(f"PERT 母版不存在：{path}（可设环境变量 WE_TPL_PERT 指定）")
    shutil.copy2(path, out)
    wb = openpyxl.load_workbook(out, data_only=False)
    ws = wb[SHEET]
    ws["D3"] = project_name
    if est_date:
        ws["D2"] = est_date
    if estimator:
        ws["H2"] = estimator
    # 清空模板自带的示例数据（D/E/F/G），H/I 公式列保留由 Excel 重算
    for r in range(DEV_START, DEV_END + 1):
        for c in (4, 5, 6, 7):
            ws.cell(r, c).value = None
    for r in PHASE_ROW.values():
        for c in (5, 6, 7):
            ws.cell(r, c).value = None
    dev_idx = DEV_START
    for r in rows:
        ph = r.get("phase") or "开发"
        if ph == "开发" and dev_idx <= DEV_END:
            # B/C 列为模板合并分区（MergedCell 只读），仅写入 D 描述与 E/F/G 三值；
            # 系统/小模块语义纳入 desc 前缀，避免破坏合并结构。
            desc = r.get("desc") or ""
            if r.get("system") and desc:
                desc = f"{r['system']}·{desc}" if desc else r["system"]
            if r.get("module") and desc:
                desc = f"{r['module']}·{desc}" if desc else r["module"]
            if desc:
                ws.cell(dev_idx, 4).value = desc
            ws.cell(dev_idx, 5).value = r.get("most_likely")
            ws.cell(dev_idx, 6).value = r.get("optimistic")
            ws.cell(dev_idx, 7).value = r.get("pessimistic")
            dev_idx += 1
        else:
            row = PHASE_ROW.get(ph)
            if row:
                ws.cell(row, 5).value = r.get("most_likely")
                ws.cell(row, 6).value = r.get("optimistic")
                ws.cell(row, 7).value = r.get("pessimistic")
    wb.calculation.fullCalcOnLoad = True   # 强制 Excel 打开重算 H/I 公式
    wb.save(out)
    return out


def rows_for_generate(con, pid):
    """从库取项目+明细，组装成 generate 需要的 rows。"""
    p = con.execute("SELECT id,name,estimator FROM projects WHERE id=?", (pid,)).fetchone()
    name = p[1] if p else "未命名项目"
    est = p[2] if p else ""
    items = con.execute(
        "SELECT phase,system,module,desc,optimistic,most_likely,pessimistic "
        "FROM pert_items WHERE project_id=? ORDER BY id", (pid,)).fetchall()
    rows = [{"phase": i[0], "system": i[1], "module": i[2], "desc": i[3],
             "optimistic": i[4], "most_likely": i[5], "pessimistic": i[6]} for i in items]
    return name, est, rows


def db_for(app_data):
    """默认 PERT 库路径（data/pert.db）。"""
    return str(Path(app_data) / "pert.db")