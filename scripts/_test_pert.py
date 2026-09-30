import sys, sqlite3, tempfile, os
sys.path.insert(0, r"d:\code\otherProjects\27_Evaluation\scripts")
import pert

# 临时库
tmp = tempfile.mkdtemp()
dbp = os.path.join(tmp, "p.db")
pert.PertDB(dbp)
con = sqlite3.connect(dbp)

# 建项目 + 明细
con.execute("INSERT INTO projects(name,estimator) VALUES('测试项目','张三')")
pid = con.execute("SELECT last_insert_rowid()").fetchone()[0]
items = [
    ("需求分析", "", "", "需求梳理", 5, 8, 12),
    ("开发", "票据系统", "扣款", "批量扣款提交", 6, 6, 7),
    ("开发", "票据系统", "止付", "止付结清", 8, 10, 12),
    ("实施", "", "", "部署调测", 30, 30, 35),
]
for ph, s, mod, d, o, m, p in items:
    con.execute("INSERT INTO pert_items(project_id,phase,system,module,desc,optimistic,most_likely,pessimistic) VALUES(?,?,?,?,?,?,?,?)",
                (pid, ph, s, mod, d, o, m, p))
con.commit()

# 计算
s = pert.summarize(con, pid)
print("=== summarize ===")
print("total_days:", s["total_days"], "ideal_pm:", s["ideal_person_month"], "actual_pm:", s["actual_person_month"])
for ph in s["phases"]:
    print("  ", ph)

# 生成
out = os.path.join(tmp, "out.xlsx")
name, est, rows = pert.rows_for_generate(con, pid)
dirout = os.path.dirname(out)
pert.generate(name, rows, out, est_date="2026-09-08", estimator=est,
              tpl=r"C:\Users\hspcadmin\Desktop\工作量评估\阜新银行-票据系统优化-工作量评估V2.0.xlsx")
print("\n=== generate ===")
print("generated:", os.path.exists(out), os.path.getsize(out) if os.path.exists(out) else 0)

import openpyxl
wb = openpyxl.load_workbook(out, data_only=True)
ws = wb["项目工作量估算表"]
print("D3:", ws["D3"].value, "| H2:", ws["H2"].value)
print("开发明细R8 E/F/G:", ws.cell(8,5).value, ws.cell(8,6).value, ws.cell(8,7).value)
print("需求R6 E/F/G:", ws.cell(6,5).value, ws.cell(6,6).value, ws.cell(6,7).value)
print("实施R32 E/F/G:", ws.cell(32,5).value, ws.cell(32,6).value, ws.cell(32,7).value)
con.close()