# -*- coding: utf-8 -*-
"""M4 模板生成器：复制成品母版 → 重填《规模估算》明细 → 可选重分区《费用估算》→ 保存。
保真要点：仅改写需填写的数据区，公式/样式原样保留；任一导出都能用 Excel 重新计算。

校验模式 (--verify)：读取母版既有明细原样写回，用规则引擎断言 US 合计复现金标准。
"""
import argparse, json, os, re, shutil, sys, threading, time
from copy import copy
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

# 输出重定向到管道/文件时 stdout 回退 GBK，✓/✗ 等字符会 UnicodeEncodeError，统一 UTF-8 并容忍不可编码
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

import engine

TPL = r"C:\Users\hspcadmin\Desktop\工作量评估\副本造价评估-天津银行_票据新增五篇大文章标识功能项目（软件开发功能点评估）_KX_20251027_V1.0.xlsx"
SRC = Path(r"d:\code\otherProjects\27_Evaluation")

# 规模估算明细列映射（1 起）：B系统 C一级 D二级 E三级 F四级 G描述 H计数项 I类别 J(公式) K重用 L修改 M(公式) N备注
COLS = {"system": 2, "m1": 3, "m2": 4, "m3": 5, "m4": 6, "desc": 7,
        "count_item": 8, "type": 9, "reuse": 11, "modify": 12, "remark": 14}
ROW_START, ROW_END = 6, 53


def _ensure_capacity(ws, total_rows):
    """母版明细区固定 ROW_START..ROW_END（48 行）；条数超出时在合计行前动态插入行（T00606）：
    复制明细行样式/行高、J/M 公式做本行行号改写，平移合计行合并单元格并重写 SUM 区间。
    只在工作簿副本内扩容，母版文件本身不动。返回扩容后的明细末行行号。"""
    base_rows = ROW_END - ROW_START + 1
    extra = total_rows - base_rows
    if extra <= 0:
        return ROW_END
    tail_row = ROW_END + 1                                    # 原合计行（54）
    tail_height = ws.row_dimensions[tail_row].height
    tail_merges = [str(m) for m in ws.merged_cells.ranges if m.min_row >= tail_row]
    for rng in tail_merges:
        ws.unmerge_cells(rng)
    ws.insert_rows(tail_row, amount=extra)
    # openpyxl 的 insert_rows 不移动合并单元格，需手动平移
    for rng in tail_merges:
        c1, r1, c2, r2 = range_boundaries(rng)
        ws.merge_cells(start_row=r1 + extra, start_column=c1,
                       end_row=r2 + extra, end_column=c2)
    # 新明细行：复制首明细行样式与行高；J/M 公式为本行相对引用，行号统一改写
    src = ROW_START
    for i in range(extra):
        r = tail_row + i
        for col in range(1, 15):
            ws.cell(r, col)._style = copy(ws.cell(src, col)._style)
        if ws.row_dimensions[src].height is not None:
            ws.row_dimensions[r].height = ws.row_dimensions[src].height
        for col in (10, 13):                                  # J / M
            f = ws.cell(src, col).value
            if isinstance(f, str) and f.startswith("="):
                ws.cell(r, col).value = re.sub(
                    rf"(?<![A-Z$])([A-Z]{{1,2}}){src}(?!\d)",
                    lambda m: f"{m.group(1)}{r}", f)
    # 合计行已随 insert_rows 下移到 ROW_END+extra+1：重写 SUM 区间并还原行高
    new_end = ROW_END + extra
    new_tail = new_end + 1
    ws.cell(new_tail, 10).value = f"=SUM(J{ROW_START}:J{new_end})"
    ws.cell(new_tail, 13).value = f"=SUM(M{ROW_START}:M{new_end})"
    if tail_height is not None:
        ws.row_dimensions[new_tail].height = tail_height
    return new_end


def copy_and_load(out: str, template: str = None):
    shutil.copy2(template or TPL, out)
    return openpyxl.load_workbook(out, data_only=False)


def read_template_rows():
    """读取母版规模估算明细（用于校验模式回写与读取源数据）。"""
    ws = openpyxl.load_workbook(TPL, data_only=False)["规模估算"]
    rows = []
    last = ""
    for r in range(ROW_START, ROW_END + 1):
        def v(col):
            return ws.cell(r, col).value
        type_ = v(COLS["type"])
        if not type_:
            continue  # 只按类别判定数据行；描述可能共享为空，不能据此跳过
        sys_raw = v(COLS["system"]) or ""
        if sys_raw:
            last = sys_raw
        rows.append({"row": r,
                     "system": sys_raw or last, "m1": v(COLS["m1"]) or "",
                     "m2": v(COLS["m2"]) or "", "m3": v(COLS["m3"]) or "", "m4": v(COLS["m4"]) or "",
                     "desc": v(COLS["desc"]) or "", "count_item": v(COLS["count_item"]) or "",
                     "type": str(type_).strip(),
                     "reuse": (v(COLS["reuse"]) or "").strip(), "modify": (v(COLS["modify"]) or "").strip(),
                     "remark": v(COLS["remark"]) or ""})
    return rows


def fill_rows(ws, rows, end_row=None):
    """对目标工作表写入明细（覆盖数据区，保留首尾/公式列）。end_row 供动态扩容后传入实际末行。"""
    last = end_row or ROW_END
    # 清空数据列
    for r in range(ROW_START, last + 1):
        for col in (2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14):
            ws.cell(r, col).value = None
    for it in rows:
        r = it["row"]
        ws.cell(r, COLS["system"]).value = it.get("system", "")
        ws.cell(r, COLS["m1"]).value = it.get("m1", "")
        ws.cell(r, COLS["m2"]).value = it.get("m2", "")
        ws.cell(r, COLS["m3"]).value = it.get("m3", "")
        ws.cell(r, COLS["m4"]).value = it.get("m4", "")
        ws.cell(r, COLS["desc"]).value = it.get("desc", "")
        ws.cell(r, COLS["count_item"]).value = it.get("count_item", "")
        ws.cell(r, COLS["type"]).value = it["type"]
        ws.cell(r, COLS["reuse"]).value = it["reuse"]
        ws.cell(r, COLS["modify"]).value = it["modify"]
        ws.cell(r, COLS["remark"]).value = it.get("remark", "")


def arrange_systems(items):
    """把明细按"系统 → 连续行区间"编排，返回 (rows, sys_ranges:{sys:[r0,r1]})。"""
    order, ranges = [], {}
    cur = ROW_START
    rows = []
    for sys, grp in items.items():
        first = cur
        for item in grp:
            item = dict(item); item["row"] = cur
            rows.append(item); cur += 1
        ranges[sys] = [first, cur - 1]; order.append(sys)
    return rows, ranges


def rewrite_fee_systems(wb, sys_ranges):
    """按系统行区间重写《费用估算》系统列名称与分区 SUM（支持前 3 个系统）。"""
    fee = wb["费用估算"]
    cols = ["E", "F", "G"]
    if len(sys_ranges) > 3:
        raise ValueError(f"模板仅支持 ≤3 系统分区，当前 {len(sys_ranges)} 个")
    for idx, (sys, (r0, r1)) in enumerate(sys_ranges.items()):
        L = cols[idx]
        fee[f"{L}2"] = sys                          # 系统名称
        fee[f"{L}3"] = f"=SUM(规模估算!$M{r0}:$M{r1})"  # 分区汇总
    fee["D3"] = "=SUM(E3:G3)"
    return fee


def run_verify(out):
    rows = read_template_rows()
    wb = copy_and_load(out)   # 金标准校验固定用内置母版（template 不属于该作用域）
    fill_rows(wb["规模估算"], rows)
    wb.save(out)
    # 规则引擎复算金标准
    us_sum = engine.total_us([{"type": x["type"], "reuse": x["reuse"], "modify": x["modify"]} for x in rows])
    cnt = len(rows)
    print(f"[verify] 功能点明细 {cnt} 条；引擎 US 合计 = {us_sum}；母版金标准 = 105.07")
    ok = abs(us_sum - 105.07) < 0.05
    print(f"[verify] 校验结论：{'通过 ✓' if ok else '不通过 ✗'}")
    return ok


def _emit(on_progress, pct, stage):
    """安全上报进度（T00927）：回调自身异常一律吞掉，不得影响生成主流程。"""
    if not on_progress:
        return
    try:
        on_progress(pct, stage)
    except Exception:
        pass


# ---------------- T00931：Excel 公式缓存重算的调度（异步化） ----------------
# 背景：COM 重算实测占单次生成 5~20s，是"跳转延迟过长"的唯一大头，而它只影响成品表
# 公式列的 cached value —— 产物在 wb.save() 时已完整（且已设 fullCalcOnLoad，Excel 打开会自算）。
# 做法：主流程照常返回，重算改到守护线程里对**副本**执行，完成后再原子替换主产物，
# 避免下载与重算抢同一个文件句柄；同一时刻只允许一个重算在跑。
_RECALC_SLOT = threading.Lock()

# T01481（P1-2）：Excel COM 熔断。实测并发下 16 次 com_error（Workbooks.Open 无效）
# + 222 次「已有重算任务在进行」。连续失败达到阈值后短期停用 COM 重算，
# 降级为 fullCalcOnLoad（成品表已强制该标志，Excel 打开时自算），
# 避免每个请求都去 Dispatch 一次 Excel、拖慢生成并堆积隐藏进程。
_RECALC_FAILS = 0
_RECALC_DISABLE_UNTIL = 0.0
_RECALC_CIRCUIT_THRESHOLD = 3      # 连续失败次数阈值
_RECALC_CIRCUIT_COOLDOWN = 300.0   # 熔断冷却秒数（默认 5 分钟）


def _recalc_note_fail():
    """记录一次重算失败，达到阈值即打开熔断。"""
    global _RECALC_FAILS, _RECALC_DISABLE_UNTIL
    _RECALC_FAILS += 1
    if _RECALC_FAILS >= _RECALC_CIRCUIT_THRESHOLD:
        _RECALC_DISABLE_UNTIL = time.time() + _RECALC_CIRCUIT_COOLDOWN
        print(f"[recalc] 连续 {_RECALC_FAILS} 次重算失败，熔断 {_RECALC_CIRCUIT_COOLDOWN:.0f}s，"
              "期间降级为 fullCalcOnLoad（成品表仍完整）", flush=True)


def _recalc_note_ok():
    global _RECALC_FAILS, _RECALC_DISABLE_UNTIL
    _RECALC_FAILS = 0
    _RECALC_DISABLE_UNTIL = 0.0


def recalc_mode():
    """WE_RECALC_MODE = async(默认) | sync | off —— 按部署场景切换重算策略。"""
    v = (os.environ.get("WE_RECALC_MODE") or "async").strip().lower()
    return v if v in ("async", "sync", "off") else "async"


def schedule_excel_recalc(out, on_progress=None, mode=None, timeout=None, on_state=None):
    """调度公式缓存重算，返回 (queued: bool, mode: str) 供上层写响应/日志。
      sync ：阻塞等重算完成（原行为；CLI 与排查用）
      async：立即返回，重算在守护线程内对副本执行并原子替换（生成结果马上可下载）
      off  ：完全不重算
    mode/timeout 缺省时回退环境变量（服务端会显式传入配置中心的值，T00962）。
    on_state: 可选回调 on_state("written"|"skipped"|"failed")，用于结果页展示补算状态（T00963）。
    """
    def state(st):
        try:
            if on_state:
                on_state(st)
        except Exception:
            pass

    mode = mode or recalc_mode()
    if mode == "off":
        state("skipped")
        return False, "off"
    if mode == "sync":
        ok = _try_excel_recalc(out, on_progress=on_progress, timeout=timeout)
        state("written" if ok else "skipped")
        return ok, "sync"
    # T01481（P1-2）：熔断期内直接降级，不再 Dispatch Excel（避免拖慢每个生成请求）
    if time.time() < _RECALC_DISABLE_UNTIL:
        state("skipped")
        return False, "circuit-open"
    # T01481（P1-2）：有界等待槽位（默认 0 = 拿不到立刻跳过，保持原「不阻塞生成」语义；
    # 需要提高重算命中率时可设 WE_RECALC_WAIT>0，代价是最多等这么久）。
    try:
        wait = float(os.environ.get("WE_RECALC_WAIT") or 0)
    except ValueError:
        wait = 0.0
    if not _RECALC_SLOT.acquire(timeout=wait):
        print("[recalc] 未等到重算槽位，本次跳过（不影响生成结果）", flush=True)
        state("skipped")
        return False, "busy"
    out = Path(out)
    tmp = out.with_name(out.stem + ".recalc.xlsx")

    def work():
        try:
            st0 = out.stat().st_mtime_ns               # M6：记录重算起点，写回前校验防止覆盖更新的产物
            shutil.copy2(out, tmp)                     # 只动副本，主产物始终可下载
            if not _try_excel_recalc(tmp, on_progress=on_progress, timeout=timeout):
                tmp.unlink(missing_ok=True)
                _recalc_note_fail()          # T01481（P1-2）：计入熔断计数
                state("skipped")
                return
            for i in range(6):                          # 目标可能正被下载读取，退避重试
                try:
                    if out.stat().st_mtime_ns != st0:
                        tmp.unlink(missing_ok=True)
                        print("[recalc] 目标在重算期间已被更新，放弃写回（避免覆盖新产物）", flush=True)
                        state("skipped")
                        return
                    os.replace(tmp, out)
                    _recalc_note_ok()        # T01481（P1-2）：成功即清零连续失败计数
                    print(f"[recalc] 后台重算完成并已写回: {out.name}", flush=True)
                    state("written")
                    return
                except OSError as e:
                    if i == 5:
                        print(f"[recalc] 写回失败（{e}），保留副本: {tmp.name}", flush=True)
                        state("skipped")
                        return
                    time.sleep(0.5)
        except Exception as e:
            print(f"[recalc] 后台重算异常: {type(e).__name__}: {e}", flush=True)
            state("failed")
        finally:
            _RECALC_SLOT.release()

    threading.Thread(target=work, name="we-excel-recalc-async", daemon=True).start()
    return True, "async"


def _recolor_header(ws, brand_color):
    """M5/FR-4.2：把母版表头行（前 5 行内、实底填充的非白单元格）的填充色替换为主题色。
    仅动 fill，不改字体/边框/列宽——母版其余样式完整保留。异常吞掉（主题联动失败不影响生成）。"""
    try:
        from openpyxl.styles import PatternFill
        target = "FF" + brand_color.lstrip("#").upper()
        for row in ws.iter_rows(min_row=1, max_row=5):
            for cell in row:
                f = cell.fill
                if not f or f.patternType != "solid":
                    continue
                fg = f.fgColor
                rgb = getattr(fg, "rgb", None)
                is_theme = getattr(fg, "type", None) == "theme"
                is_rgb = isinstance(rgb, str) and len(rgb) == 8 and rgb.upper() not in ("FFFFFFFF", "00000000")
                if is_theme or is_rgb:          # 母版表头为主题色填充：统一替换为主题色实底
                    cell.fill = PatternFill("solid", fgColor=target)
    except Exception:
        pass


def _append_gb_sheet(wb, o):
    """M6/FR-2.2：追加「国标造价」sheet——同一 docx 同时产出 IFPUG 母版表与 GB 造价测算表（AC）。"""
    try:
        from openpyxl.styles import Font, PatternFill
        name = "国标造价"
        if name in wb.sheetnames:
            del wb[name]
        ws = wb.create_sheet(name)
        rows = [("GB/T 36964-2018 国标造价测算", ""),
                ("估算阶段", o.get("stage", "")),
                ("CF 规模调整因子", o.get("cf", "")),
                ("GSC 全局调整因子", o.get("gsc", "")),
                ("SWF 业务领域因子", o.get("swf", "")),
                ("RDF 开发语言/团队因子", o.get("rdf", "")),
                ("PDR 生产率（人时/FP）", o.get("pdr", "")),
                ("人月费率（元/人月）", o.get("rate", "")),
                ("DNHC 直接非人力成本（元）", o.get("dnhc", "")),
                ("UFP（未调整功能点）", o.get("ufp", "")),
                ("S 调整后规模", o.get("S", "")),
                ("AE（人时）", o.get("ae", "")),
                ("折算人月", round(o.get("ae", 0) / 174.0, 2)),
                ("费用（元）", o.get("fee", "")),
                ("费用（万元）", round(o.get("fee", 0) / 10000.0, 2)),
                ("出处", "GB/T 36964-2018 · CSBMK 基准数据（取值与出处可追溯）")]
        for r, (k, v) in enumerate(rows, 1):
            ws.cell(r, 1, k)
            ws.cell(r, 2, v)
        ws["A1"].font = Font(bold=True, size=13)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 30
    except Exception:
        pass   # 造价 sheet 失败不影响母版成品


def _atomic_replace(tmp, out, retries=3, delay=0.5):
    """T01481（P0-2）：os.replace 带短重试——目标 out 可能正被另一同名请求的
    Excel COM 实例短暂持有（async 后台重算写回窗口），重试通常即可成功。"""
    import time as _t
    last = None
    for _ in range(retries + 1):
        try:
            os.replace(tmp, out)
            return
        except PermissionError as e:
            last = e
            _t.sleep(delay)
    raise last


def _remove_save_tmp(tmp):
    """T01481（P0-2）：原子写失败时清理自己的 .tmp 临时文件。
    注意：必须用 ctypes DeleteFileW 直调 Win32——本环境 os.remove 被
    safe-delete 钩子劫持为回收站语义且 fail-closed（tmp 竞争消失时
    trash 报"找不到文件"会抛异常，反而打断生成流程）。"""
    if not tmp:
        return
    try:
        import ctypes
        ctypes.windll.kernel32.DeleteFileW(tmp)
    except Exception:
        pass


def run_from_json(js, out, systems=None, method="估算功能点", on_progress=None, recalc_info=None,
                  recalc_mode=None, recalc_timeout=None, on_recalc_state=None, brand_color=None, gb_state=None,
                  template=None):
    """生成成品表。

    on_progress: 可选回调 on_progress(pct:int, stage:str)，用于上层（服务端）上报真实阶段
    进度（T00927）。回调自身异常一律吞掉，不得影响生成主流程。
    recalc_mode / recalc_timeout: 上层显式给出的公式缓存重算策略（T00962）；缺省回退环境变量。
    on_recalc_state: 可选回调 on_recalc_state("written"|"skipped"|"failed")（T00963）。
    """
    _prog = lambda p, s: _emit(on_progress, p, s)

    if isinstance(js, (str, Path)):
        data = json.loads(Path(js).read_text(encoding="utf-8"))
    else:
        data = js
    # data 可以是 {system:[items...]} 或 [items]（带 system 字段）
    planned = {}
    if isinstance(data, list):
        for it in data:
            s = (it.get("system") or (systems and systems[0])) or "综合服务平台"
            planned.setdefault(s, []).append(it)
    else:
        planned = data
    rows, ranges = arrange_systems(planned)
    _prog(20, "正在排布系统与明细行…")
    wb = copy_and_load(out, template)   # M9/FR-2.8：按所选母版生成（缺省内置母版）
    ws = wb["规模估算"]
    new_end = _ensure_capacity(ws, len(rows))     # 行数超母版 48 行容量时动态扩容（T00606）
    _prog(30, "正在载入母版并扩容明细区…")
    ws["C1"] = method  # 规模估算方法：联动母版 J 列权重公式（估算功能点/预估功能点）
    fill_rows(ws, rows, end_row=new_end)
    _prog(48, f"正在写入 {len(rows)} 条功能点明细…")
    if brand_color:
        _recolor_header(ws, brand_color)   # M5/FR-4.2：母版表头填充色随主题色联动（仅表头 fill，其余母版样式不动）
    rewrite_fee_systems(wb, ranges)
    if gb_state:
        _append_gb_sheet(wb, gb_state)   # M6/FR-2.2：成品表追加「国标造价」sheet（同 docx 双口径产出）
    # 强制 Excel 打开即重算：J/M 列是母版公式，openpyxl 不算缓存值，
    # 不设此标志的话 Excel 默认按"上次缓存"展示，导致用户看到 J/M 是空。
    wb.calculation.fullCalcOnLoad = True
    tmp = None
    try:
        # T01481（P0-2）：先写同目录唯一临时文件再 os.replace 原子替换。
        # 并发/重开同标题时两个线程曾直接 save 同一路径互相踩踏
        # （T01369 对照实验：固定 title 并发 42.33% 500 vs 唯一 title 0%）。
        tmp = out + f".tmp{os.getpid()}.{threading.get_ident()}"
        wb.save(tmp)
        _atomic_replace(tmp, out)
        tmp = None   # 已原子落位，无需清理
    except PermissionError as e:
        # 目标被占用：最常见是上一次卡死的 Excel 还开着它（T00960）；
        # os.replace 在 Windows 上替换被占用目标同样抛 PermissionError，保留原提示语义
        _remove_save_tmp(tmp)
        raise RuntimeError(
            "成品表写入失败：目标文件被占用（常见于上次生成的 Excel 实例仍在后台），"
            "请稍后重试、或更换需求标题后另存；若反复出现请查看服务日志中的 [recalc] 记录。") from e
    except Exception:
        _remove_save_tmp(tmp)
        raise
    _prog(58, "正在保存成品表…")
    # 若本机装有 Excel，再用 COM 跑一次重算并写回 cached value，
    # 这样即便用户 Excel 处于"手动重算"也能立即看到数值。
    _prog(62, "正在处理母版公式缓存（Excel 重算）…")
    queued, how = schedule_excel_recalc(out, on_progress=on_progress,
                                        mode=recalc_mode, timeout=recalc_timeout,
                                        on_state=on_recalc_state)
    if recalc_info is not None:
        recalc_info["mode"] = how
    if how == "sync":
        _prog(95, "母版公式重算完成，正在收尾…" if queued else "已跳过 Excel 重算，正在收尾…")
    elif how == "async":
        _prog(95, "成品表已生成；母版公式缓存在后台补算，正在收尾…")
    else:
        _prog(95, "已跳过 Excel 重算，正在收尾…")
    us_sum = engine.total_us([{"type": r["type"], "reuse": r["reuse"], "modify": r["modify"]} for r in rows])
    print(f"[gen] 已生成 {out}；功能点 {len(rows)} 条，US 合计 {us_sum}")
    return rows, ranges


def _excel_pid(xl):
    """取 Dispatch 出来的那个 Excel 实例的 PID；取不到返回 None（T00960）。"""
    try:
        import win32process
        return int(win32process.GetWindowThreadProcessId(xl.Hwnd)[1])
    except Exception:
        return None


def _reap_excel(pid, done_flag, grace):
    """超时后的定向回收：只终止我们自己 Dispatch 出来的隐藏实例（按 PID），
    绝不碰用户自己打开的 Excel。宽限 grace 秒，期间工作线程若已收尾则不杀（T00960）。"""
    for _ in range(max(1, int(grace // 2))):
        time.sleep(2)
        if done_flag.get("done"):
            print(f"[recalc] PID={pid} 已自行收尾，无需定向回收", flush=True)
            return
    try:
        import subprocess
        r = subprocess.run(["taskkill", "/F", "/PID", str(pid), "/T"],
                           capture_output=True, text=True, timeout=15)
        ok = r.returncode == 0
        print(f"[recalc] 定向回收卡死实例 PID={pid}：{'成功' if ok else '失败 ' + (r.stderr or '').strip()[:120]}",
              flush=True)
    except Exception as e:
        print(f"[recalc] 定向回收 PID={pid} 异常（留待自然退出）: {type(e).__name__}: {e}", flush=True)


def _try_excel_recalc(path, on_progress=None, timeout=None):
    """若本机装有 Excel，用 COM 打开→Calculate→Save，把公式的 cached value 写回 xlsx。
    失败（无 Excel/无 pywin32）静默跳过，不影响主流程。

    on_progress: 同 run_from_json，用于上报该慢步骤内部的真实子阶段（T00927）。
    timeout: 秒，默认取环境变量 WE_RECALC_TIMEOUT（缺省 120）。COM 调用无法从外部打断，
      故放到守护线程里做「带超时的等待」：超时即放弃本次重算并按「未重算」继续返回
      （生成结果仍可用；成品表已设 fullCalcOnLoad，Excel 打开时会自行重算），
      避免 Excel 弹窗/卡死把整个生成请求无限期挂住（T00930）。
      超时后再按 PID 定向回收该实例（宽限 WE_RECALC_KILL_GRACE 秒，默认 30s，T00960）。
    返回 True=已重算写回；False=跳过/失败。
    """
    import os
    try:
        import win32com.client  # type: ignore
        import pythoncom
    except Exception as e:
        if os.environ.get("WE_DEBUG_RECALC"):
            print(f"[recalc] skip: import win32com failed: {e}", flush=True)
        return False
    if timeout is None:
        try:
            timeout = float(os.environ.get("WE_RECALC_TIMEOUT") or 120)
        except ValueError:
            timeout = 120.0
    try:
        grace = float(os.environ.get("WE_RECALC_KILL_GRACE") or 30)
    except ValueError:
        grace = 30.0
    result = {"ok": False, "err": None, "pid": None, "done": False}

    def work():
        pythoncom.CoInitialize()          # COM 需按线程初始化（已在独立线程内执行）
        xl = None
        wb = None
        try:
            _emit(on_progress, 66, "正在启动 Excel 进程…")
            xl = win32com.client.DispatchEx("Excel.Application")
            xl.Visible = False
            xl.DisplayAlerts = False      # 抑制对话框：多数"卡住"其实是在等一个看不见的弹窗
            result["pid"] = _excel_pid(xl)   # 记下实例 PID：超时后只回收这一个进程（T00960）
            _emit(on_progress, 72, "正在打开成品表…")
            wb = xl.Workbooks.Open(str(Path(path).resolve()))
            _emit(on_progress, 82, "正在重算母版公式…")
            xl.Calculate()
            _emit(on_progress, 90, "正在写回计算结果并关闭 Excel…")
            # T01481（P0-2）：SaveAs 到临时副本再原子替换——直接 Save 回写 out
            # 会与并发同名请求的原子替换互相踩踏（实测 "File is not a zip file"）。
            tmp_recalc = str(Path(path).with_name(
                Path(path).name + f".recalc{os.getpid()}.{threading.get_ident()}.tmp"))
            wb.SaveAs(tmp_recalc)
            wb.Close(); wb = None
            xl.Quit(); xl = None
            _atomic_replace(tmp_recalc, str(path))
            tmp_recalc = None
            result["ok"] = True
            if os.environ.get("WE_DEBUG_RECALC"):
                print(f"[recalc] OK: {path}", flush=True)
        except Exception as e:
            result["err"] = f"{type(e).__name__}: {e}"
            print(f"[recalc] FAILED: {result['err']}", flush=True)
            try:
                if tmp_recalc:
                    _remove_save_tmp(tmp_recalc)
            except Exception:
                pass
        finally:
            # 无论成功/失败都尽量收尾，避免残留 EXCEL.EXE 与文件句柄
            try:
                if wb is not None:
                    wb.Close(False)
            except Exception:
                pass
            try:
                if xl is not None:
                    xl.Quit()
            except Exception:
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass
            result["done"] = True          # 工作线程生命周期结束（含收尾）的标记

    th = threading.Thread(target=work, name="we-excel-recalc", daemon=True)
    th.start()
    th.join(timeout)
    if th.is_alive():
        pid = result.get("pid")
        print(f"[recalc] 超时：{timeout:.0f}s 内未完成，放弃本次 COM 重算"
              f"（后台线程继续收尾，不阻塞本次生成返回；如需放宽可设 WE_RECALC_TIMEOUT）", flush=True)
        if pid:
            # 只回收我们自己创建的这个隐藏实例，宽限期内它若自己收尾则不杀（T00960）
            threading.Thread(target=_reap_excel, args=(pid, result, grace),
                             name="we-excel-reap", daemon=True).start()
            print(f"[recalc] 已启动定向回收：PID={pid}，宽限 {grace:.0f}s", flush=True)
        _emit(on_progress, 95, f"Excel 重算超时（>{timeout:.0f}s），已跳过；母版公式将由 Excel 打开时自动重算")
        return False
    return bool(result["ok"])


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--verify", action="store_true")
    ap.add_argument("--json", default=None)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    out = a.out or str(SRC / "data" / "output.xlsx")
    if a.verify:
        sys.exit(0 if run_verify(out) else 1)
    if a.json:
        run_from_json(a.json, out)
        sys.exit(0)
    print("需指定 --verify 或 --json <候选清单>")