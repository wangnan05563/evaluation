# -*- coding: utf-8 -*-
"""探测任意需求文档：段落/表格规模、含数据表格、功能需求清单表定位。用法: python probe_doc.py <docx路径>"""
import sys, re
from docx import Document
from docx.oxml.ns import qn

path = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\hspcadmin\Desktop\工作量评估\HNNS-EB-STD-REQ-002-票据业务系统反洗钱工作对接-日常优化类(公司金融部)(20260831).docx"
doc = Document(path)

np = len(doc.paragraphs)
ne = sum(1 for p in doc.paragraphs if p.text.strip())
print(f"段落 {np} / 非空 {ne} ; 表格 {len(doc.tables)}")

# 含数据表格概览
print("\n含数据表格(前20):")
shown = 0
for i, tb in enumerate(doc.tables[:40], 1):
    head = " | ".join((c.text or "").strip()[:12] for c in tb.rows[0].cells)
    data_rows = sum(1 for r in tb.rows[1:] if any((c.text or "").strip() for c in r.cells))
    if data_rows:
        print(f"  表{i} {len(tb.rows)}行 head[{head}] 数据行{data_rows}")
        shown += 1
        if shown >= 20: break

# 定位功能需求清单表（字段含 模块/功能/子功能/新增）
print("\n定位功能清单表:")
for i, tb in enumerate(doc.tables, 1):
    head = " ".join((c.text or "") for c in tb.rows[0].cells)
    if any(k in head for k in ("模块", "功能", "子功能")):
        data_rows = sum(1 for r in tb.rows[1:] if any((c.text or "").strip() for c in r.cells))
        print(f"  表{i} head[:50]={head[:50]!r} 数据行={data_rows}")
        if data_rows:
            for r in tb.rows[1:4]:
                print("      " + " | ".join((c.text or "").strip()[:18] for c in r.cells))

# 标题结构（前若干非空标题）
print("\n标题(前20非空):")
n = 0
for p in doc.paragraphs:
    t = p.text.strip()
    st = (p.style.name or "").lower() if p.style else ""
    if t and st.startswith("heading"):
        print(f"  [{st}] {t[:50]}")
        n += 1
        if n >= 20: break